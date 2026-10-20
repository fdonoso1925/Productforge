"""Planillas funcionales: XLSX (openpyxl) con fórmulas reales + un CSV por pestaña.

La IA describe las columnas calculadas como operaciones (`Computation`); aquí se
traducen a fórmulas de Excel. Las listas de opciones viven en una hoja oculta, así no
hay límite de 255 caracteres en la validación.
"""

import csv
import io
import re
from datetime import datetime
from typing import ClassVar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from app.content.schemas import Format
from app.content.schemas.formats import Column, Computation, SheetContent, SheetTab
from app.generators.base import FileArtifact, RenderContext, artifact, register, slugify
from app.templates_engine.model import OfficeTheme, Template

LISTS_SHEET = "_listas"
CURRENCY_FORMAT = {"BRL": '"R$" #,##0.00', "USD": '"US$" #,##0.00'}
TOTAL_FUNCTIONS = {
    "sum": "SUM",
    "average": "AVERAGE",
    "count": "COUNTA",
    "max": "MAX",
    "min": "MIN",
}
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def fill(hex_color: str) -> PatternFill:
    color = hex_color.lstrip("#").upper()
    return PatternFill("solid", start_color=color, end_color=color)


def formula(comp: Computation, refs: dict[str, str], row: int) -> str:
    cells = [f"{refs[name]}{row}" for name in comp.operands]
    match comp.op:
        case "sum":
            return f"=SUM({','.join(cells)})"
        case "subtract":
            return "=" + "-".join(cells)
        case "multiply":
            return "=" + "*".join(cells)
        case "divide" | "percent_of":
            return f"=IFERROR({cells[0]}/{'/'.join(cells[1:])},0)"


def parse_number(raw: str, percent: bool = False) -> float | None:
    text = raw.strip().replace(" ", "").replace(" ", "")
    is_percent = text.endswith("%")
    text = re.sub(r"[^\d,.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        text = text.replace(thousands, "").replace(decimal, ".")
    elif "," in text:
        whole, _, frac = text.rpartition(",")
        text = f"{whole.replace(',', '')}.{frac}" if len(frac) <= 2 else text.replace(",", "")
    try:
        value = float(text)
    except ValueError:
        return None
    if percent and (is_percent or value > 1):
        value /= 100
    return value


def parse_date(raw: str) -> datetime | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw.strip(), fmt)
        except ValueError:
            continue
    return None


def cell_value(column: Column, raw: str, labels: dict[str, str]):
    if not raw.strip():
        return None
    match column.type:
        case "number" | "currency":
            return parse_number(raw)
        case "percent":
            return parse_number(raw, percent=True)
        case "date":
            return parse_date(raw) or raw
        case "checkbox":
            yes = raw.strip().lower() in {"sí", "si", "sim", "yes", "x", "true", "1", "✔"}
            return labels["yes"] if yes else labels["no"]
        case _:
            return raw


def number_format(column: Column, currency: str) -> str | None:
    if column.computed and column.computed.op == "percent_of":
        return "0.0%"
    return {
        "currency": CURRENCY_FORMAT[currency],
        "percent": "0.0%",
        "number": "#,##0.##",
        "date": "DD/MM/YYYY",
    }.get(column.type)


def safe_sheet_title(name: str, used: set[str]) -> str:
    base = INVALID_SHEET_CHARS.sub(" ", name).strip()[:31] or "Hoja"
    title, n = base, 2
    while title.lower() in used:
        title = f"{base[:28]} {n}"
        n += 1
    used.add(title.lower())
    return title


class _Lists:
    """Hoja oculta con las opciones de cada columna `choice` / `checkbox`."""

    def __init__(self, wb: Workbook) -> None:
        self.ws = wb.create_sheet(LISTS_SHEET)
        self.ws.sheet_state = "hidden"
        self.next_col = 1

    def add(self, options: list[str]) -> str:
        col = get_column_letter(self.next_col)
        for i, option in enumerate(options, start=1):
            self.ws[f"{col}{i}"] = option
        self.next_col += 1
        return f"{quote_sheetname(LISTS_SHEET)}!${col}$1:${col}${len(options)}"


def write_tab(wb, tab: SheetTab, title: str, theme: OfficeTheme, ctx: RenderContext, lists):
    ws = wb.create_sheet(title)
    labels = ctx.labels
    refs = {c.header: get_column_letter(i) for i, c in enumerate(tab.columns, start=1)}
    first, last = 2, 1 + len(tab.sample_rows) + tab.empty_rows
    header_font = Font(name=theme.heading_font, bold=True, color=theme.on_primary.lstrip("#"))
    body_font = Font(name=theme.body_font)
    computed_fill = fill(theme.surface)

    for col_idx, column in enumerate(tab.columns, start=1):
        letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx, value=column.header)
        header.font, header.fill = header_font, fill(theme.primary)
        header.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = max(12, min(40, len(column.header) + 6))
        fmt = number_format(column, ctx.currency)

        for row in range(first, last + 1):
            cell = ws.cell(row=row, column=col_idx)
            sample_index = row - first
            if column.computed:
                cell.value = formula(column.computed, refs, row)
                cell.fill = computed_fill
            elif sample_index < len(tab.sample_rows):
                cell.value = cell_value(column, tab.sample_rows[sample_index][col_idx - 1], labels)
            cell.font = body_font
            if fmt:
                cell.number_format = fmt

        options = column.choices if column.type == "choice" else None
        if column.type == "checkbox":
            options = [labels["yes"], labels["no"]]
        if options:
            clean = [o.replace('"', "'") for o in options]
            validation = DataValidation(
                type="list", formula1=f"={lists.add(clean)}", allow_blank=True
            )
            ws.add_data_validation(validation)
            validation.add(f"{letter}{first}:{letter}{last}")

    if any(c.total for c in tab.columns):
        total_row = last + 1
        border = Border(top=Side(style="medium", color=theme.primary.lstrip("#")))
        ws.cell(row=total_row, column=1, value=labels["total"]).font = Font(
            name=theme.body_font, bold=True
        )
        for col_idx, column in enumerate(tab.columns, start=1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.border = border
            if column.total:
                letter = get_column_letter(col_idx)
                fn = TOTAL_FUNCTIONS[column.total]
                cell.value = f"={fn}({letter}{first}:{letter}{last})"
                cell.font = Font(name=theme.body_font, bold=True)
                if fmt := number_format(column, ctx.currency):
                    cell.number_format = fmt

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(tab.columns))}{last}"
    ws.sheet_properties.tabColor = theme.accent.lstrip("#")
    return ws


def build_xlsx(content: SheetContent, theme: OfficeTheme, ctx: RenderContext) -> tuple[bytes, list]:
    wb = Workbook()
    readme = wb.active
    used = {LISTS_SHEET}
    readme.title = safe_sheet_title(ctx.labels["how_to_use"], used)
    readme["A1"] = content.title
    readme["A1"].font = Font(name=theme.heading_font, size=18, bold=True, color=theme.primary[1:])
    readme["A2"] = content.purpose
    row = 4
    for line in content.how_to_use:
        readme.cell(row=row, column=1, value=f"• {line}")
        row += 1
    row += 1
    readme.column_dimensions["A"].width = 100

    lists = _Lists(wb)
    sheets = []
    for tab in content.tabs:
        title = safe_sheet_title(tab.name, used)
        readme.cell(row=row, column=1, value=f"{title}: {tab.instructions}").font = Font(
            name=theme.body_font, bold=True
        )
        row += 1
        sheets.append((tab, write_tab(wb, tab, title, theme, ctx, lists)))
    wb.move_sheet(LISTS_SHEET, offset=len(wb.sheetnames))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), sheets


def tab_csv(ws) -> bytes:
    """CSV con las fórmulas como texto: Google Sheets las evalúa al importar."""
    out = io.StringIO()
    writer = csv.writer(out)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(
            [
                v.strftime("%Y-%m-%d") if isinstance(v, datetime) else ("" if v is None else v)
                for v in row
            ]
        )
    return out.getvalue().encode("utf-8-sig")


@register
class SheetGenerator:
    format: ClassVar[Format] = Format.SHEET
    outputs: ClassVar[tuple[str, ...]] = ("xlsx", "csv")

    def generate(
        self, content: SheetContent, template: Template, ctx: RenderContext
    ) -> list[FileArtifact]:
        xlsx, sheets = build_xlsx(content, template.office_theme(), ctx)
        files = [artifact(ctx, "xlsx", xlsx)]
        files += [artifact(ctx, "csv", tab_csv(ws), suffix=slugify(tab.name)) for tab, ws in sheets]
        return files
