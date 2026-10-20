"""DOCX y PPTX editables. La plantilla solo aporta colores y tipografías (§5)."""

import io
import re
from typing import ClassVar

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches
from pptx.util import Pt as PptxPt

from app.content.schemas import Format
from app.content.schemas.formats import EditableTemplateContent, SwipeFileContent
from app.generators.base import FileArtifact, RenderContext, artifact, register
from app.templates_engine.model import OfficeTheme, Template

PLACEHOLDER = re.compile(r"(\[[^\]]+\])")


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


# ─────────────────────────── DOCX ───────────────────────────


def new_document(theme: OfficeTheme):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = theme.body_font
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(theme.text)
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = theme.heading_font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(theme.primary if name != "Heading 3" else theme.secondary)
    return doc


def add_text(paragraph, text: str, bold: bool = False) -> None:
    """Agrega texto resaltando en amarillo cada [MARCADOR] que el comprador debe reemplazar."""
    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold
        if PLACEHOLDER.fullmatch(part):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True


def shade(paragraph, hex_color: str) -> None:
    props = paragraph._p.get_or_add_pPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_color.lstrip("#"))
    props.append(fill)


def to_bytes(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def editable_docx(content: EditableTemplateContent, theme: OfficeTheme) -> bytes:
    doc = new_document(theme)
    doc.add_paragraph(content.title, style="Title")
    for line in content.instructions:
        add_text(doc.add_paragraph(style="List Bullet"), line)
    for document in content.documents:
        doc.add_page_break()
        doc.add_paragraph(document.name, style="Heading 1")
        purpose = doc.add_paragraph()
        purpose.add_run(document.purpose).italic = True
        for block in document.blocks:
            if block.kind == "heading":
                doc.add_paragraph(block.text, style="Heading 2")
            elif block.kind == "bullets":
                if block.text:
                    doc.add_paragraph(block.text, style="Heading 3")
                for item in block.items:
                    add_text(doc.add_paragraph(style="List Bullet"), item)
            else:
                add_text(doc.add_paragraph(), block.text, bold=block.kind == "placeholder")
    return to_bytes(doc)


def swipe_file_docx(content: SwipeFileContent, theme: OfficeTheme, ctx: RenderContext) -> bytes:
    labels = ctx.labels
    doc = new_document(theme)
    doc.add_paragraph(content.title, style="Title")
    doc.add_paragraph(content.introduction)
    for category in content.categories:
        doc.add_paragraph(category.name, style="Heading 1")
        when = doc.add_paragraph()
        when.add_run(f"{labels['when_to_use']}: ").bold = True
        when.add_run(category.when_to_use).italic = True
        for i, entry in enumerate(category.entries, start=1):
            doc.add_paragraph(f"{i}. {entry.label}", style="Heading 3")
            for line in entry.text.splitlines() or [""]:
                paragraph = doc.add_paragraph()
                add_text(paragraph, line)
                shade(paragraph, theme.surface)
            if entry.notes:
                notes = doc.add_paragraph()
                notes.add_run(f"{labels['notes']}: ").bold = True
                notes.add_run(entry.notes)
    return to_bytes(doc)


# ─────────────────────────── PPTX ───────────────────────────


def _style_runs(text_frame, font: str, size: int, color: str, bold: bool = False) -> None:
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = PptxPt(size)
            run.font.bold = bold
            run.font.color.rgb = PptxRGB.from_string(color.lstrip("#").upper())


def _accent_bar(slide, prs, theme: OfficeTheme) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = PptxRGB.from_string(theme.primary.lstrip("#").upper())
    bar.line.fill.background()


def editable_pptx(content: EditableTemplateContent, theme: OfficeTheme) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = content.title
    cover.placeholders[1].text = "\n".join(content.instructions)
    _style_runs(cover.shapes.title.text_frame, theme.heading_font, 40, theme.primary, bold=True)
    _style_runs(cover.placeholders[1].text_frame, theme.body_font, 18, theme.muted)
    _accent_bar(cover, prs, theme)

    for item in content.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = item.title
        body = slide.placeholders[1].text_frame
        body.text = item.bullets[0] if item.bullets else ""
        for bullet in item.bullets[1:]:
            body.add_paragraph().text = bullet
        _style_runs(slide.shapes.title.text_frame, theme.heading_font, 32, theme.primary, True)
        _style_runs(body, theme.body_font, 20, theme.text)
        _accent_bar(slide, prs, theme)
        if item.speaker_notes:
            slide.notes_slide.notes_text_frame.text = item.speaker_notes

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


@register
class EditableTemplate:
    format: ClassVar[Format] = Format.EDITABLE_TEMPLATE
    outputs: ClassVar[tuple[str, ...]] = ("docx", "pptx")

    def generate(
        self, content: EditableTemplateContent, template: Template, ctx: RenderContext
    ) -> list[FileArtifact]:
        theme = template.office_theme()
        if content.output == "pptx":
            return [artifact(ctx, "pptx", editable_pptx(content, theme))]
        return [artifact(ctx, "docx", editable_docx(content, theme))]
