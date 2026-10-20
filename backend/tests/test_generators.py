import csv
import io
import json
import shutil
import subprocess
import xml.dom.minidom
import zipfile
from pathlib import Path

import pytest
from docx import Document
from ebooklib import epub
from openpyxl import load_workbook
from pptx import Presentation

from app.config import Settings
from app.content.schemas import CommercialSheet, Format
from app.content.schemas.formats import GuideContent
from app.generators import REGISTRY, RenderContext, generate_files
from app.generators.base import slugify
from app.generators.commercial import commercial_files
from app.generators.render import pdf_page_count
from app.generators.samples import SAMPLES, SHEET, TEMPLATE_PPTX
from app.generators.sheet import parse_number
from app.templates_engine.fonts import font_path
from app.templates_engine.model import list_templates, load_template

TEMPLATES_DIR = Settings(_env_file=None).templates_dir
TEMPLATES = list_templates(TEMPLATES_DIR)
DEFAULT = load_template(TEMPLATES_DIR, "minimal-clean")


def ctx(market: str = "es-LATAM", slug: str = "pieza") -> RenderContext:
    return RenderContext(market=market, role="order_bump", slug=slug, product_name="Producto")


def assert_opens(file, tmp_path: Path) -> None:
    data = file.content
    match file.file_type:
        case "pdf":
            assert pdf_page_count(data) >= 1
        case "docx":
            assert Document(io.BytesIO(data)).paragraphs
        case "pptx":
            assert len(Presentation(io.BytesIO(data)).slides) >= 1
        case "xlsx":
            assert load_workbook(io.BytesIO(data)).sheetnames
        case "csv":
            assert list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        case "svg" | "opml" | "mm" | "drawio":
            xml.dom.minidom.parseString(data)
        case "xmind":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                sheet = json.loads(z.read("content.json"))
            assert sheet[0]["rootTopic"]["title"]
        case "epub":
            path = tmp_path / file.filename
            path.write_bytes(data)
            assert len(list(epub.read_epub(str(path)).get_items())) > 3
        case "png":
            assert data.startswith(b"\x89PNG")
        case "md" | "mmd":
            assert data.decode().strip()
        case other:
            pytest.fail(f"tipo sin verificación: {other}")


def test_every_format_has_a_generator() -> None:
    assert set(REGISTRY) == set(Format)


@pytest.mark.parametrize("fmt", list(Format))
def test_generator_outputs_open(fmt: Format, tmp_path: Path) -> None:
    files = generate_files(fmt, SAMPLES[fmt], DEFAULT, ctx(slug=fmt.value))
    assert files
    for file in files:
        assert file.file_type in REGISTRY[fmt].outputs
        assert file.filename.startswith(fmt.value)
        assert_opens(file, tmp_path)


def test_editable_template_pptx(tmp_path: Path) -> None:
    [file] = generate_files(Format.EDITABLE_TEMPLATE, TEMPLATE_PPTX, DEFAULT, ctx())
    assert file.file_type == "pptx"
    prs = Presentation(io.BytesIO(file.content))
    assert len(prs.slides) == 1 + len(TEMPLATE_PPTX.slides)
    assert prs.slides[1].notes_slide.notes_text_frame.text


def test_docx_highlights_placeholders() -> None:
    [file] = generate_files(
        Format.EDITABLE_TEMPLATE, SAMPLES[Format.EDITABLE_TEMPLATE], DEFAULT, ctx()
    )
    runs = [r for p in Document(io.BytesIO(file.content)).paragraphs for r in p.runs]
    highlighted = [r.text for r in runs if r.font.highlight_color is not None]
    assert "[NOMBRE COMPLETO]" in highlighted


@pytest.mark.parametrize("template", TEMPLATES, ids=lambda t: t.id)
def test_every_template_renders_guide_with_toc(template) -> None:
    [file] = generate_files(Format.PDF_GUIDE, SAMPLES[Format.PDF_GUIDE], template, ctx())
    # portada + índice + intro + 2 capítulos + conclusión
    assert pdf_page_count(file.content) == 6


@pytest.mark.parametrize("template", TEMPLATES, ids=lambda t: t.id)
def test_template_assets_present(template) -> None:
    for spec in (template.fonts.heading, template.fonts.body):
        for weight in spec.weights:
            assert font_path(TEMPLATES_DIR / "_fonts", spec.family, weight).exists()
    assert (template.base_dir / "preview.png").exists()


def test_at_least_ten_templates() -> None:
    assert len(TEMPLATES) >= 10


def test_html_escapes_content() -> None:
    evil = SAMPLES[Format.PDF_GUIDE].model_copy(update={"title": "<script>alert(1)</script>"})
    html = DEFAULT.render_html(
        "guide",
        c=evil,
        ctx=ctx(),
        cover={"title": evil.title, "subtitle": "", "promise": "", "product_name": "x"},
    )
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


def test_labels_follow_market() -> None:
    pt = DEFAULT.render_html(
        "guide",
        c=SAMPLES[Format.PDF_GUIDE],
        ctx=ctx("pt-BR"),
        cover={"title": "t", "subtitle": "", "promise": "", "product_name": "x"},
    )
    assert "Sumário" in pt and "Dica" in pt


# ─────────────────────────── Planilla ───────────────────────────


def sheet_workbook():
    files = generate_files(Format.SHEET, SHEET, DEFAULT, ctx())
    xlsx = next(f for f in files if f.file_type == "xlsx")
    return load_workbook(io.BytesIO(xlsx.content)), files


def test_sheet_formulas_validations_and_totals() -> None:
    wb, _ = sheet_workbook()
    ws = wb["Deudas"]
    headers = [c.value for c in ws[1]]
    cost = headers.index("Costo mensual") + 1
    assert ws.cell(row=2, column=cost).value == "=B2*C2"
    last = 1 + len(SHEET.tabs[0].sample_rows) + SHEET.tabs[0].empty_rows
    assert ws.cell(row=last + 1, column=2).value == f"=SUM(B2:B{last})"
    assert ws.cell(row=2, column=2).value == 350
    assert ws.cell(row=3, column=2).value == pytest.approx(120.5)
    assert ws.cell(row=2, column=3).value == pytest.approx(0.08)
    assert "US$" in ws.cell(row=2, column=2).number_format
    assert ws.data_validations.dataValidation
    assert wb["_listas"].sheet_state == "hidden"
    assert wb["Pagos"]["D2"].value == "=C2-B2"


def test_sheet_csv_per_tab_keeps_formulas() -> None:
    _, files = sheet_workbook()
    csvs = [f for f in files if f.file_type == "csv"]
    assert {f.filename for f in csvs} == {"pieza-deudas.csv", "pieza-pagos.csv"}
    rows = list(csv.reader(io.StringIO(csvs[0].content.decode("utf-8-sig"))))
    assert rows[1][3] == "=B2*C2"


def test_sheet_brl_currency_format() -> None:
    files = generate_files(Format.SHEET, SHEET, DEFAULT, ctx("pt-BR"))
    wb = load_workbook(io.BytesIO(files[0].content))
    assert "R$" in wb["Deudas"]["B2"].number_format


@pytest.mark.parametrize(
    ("raw", "percent", "expected"),
    [
        ("1.234,56", False, 1234.56),
        ("1,234.56", False, 1234.56),
        ("120,50", False, 120.5),
        ("R$ 19", False, 19),
        ("8%", True, 0.08),
        ("6.5", True, 0.065),
        ("0.2", True, 0.2),
        ("abc", False, None),
    ],
)
def test_parse_number(raw, percent, expected) -> None:
    assert parse_number(raw, percent) == (pytest.approx(expected) if expected else None)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice no instalado")
def test_sheet_recalculates_in_libreoffice(tmp_path: Path) -> None:
    wb, files = sheet_workbook()
    src = tmp_path / "planilla.xlsx"
    src.write_bytes(files[0].content)
    out = tmp_path / "out"
    subprocess.run(  # noqa: S603
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True,
        timeout=120,
    )
    values = load_workbook(out / "planilla.xlsx", data_only=True)["Deudas"]
    assert values["D2"].value == pytest.approx(28.0)


# ─────────────────────────── Otros ───────────────────────────


def test_mind_map_editable_formats_keep_structure() -> None:
    files = {
        f.file_type: f
        for f in generate_files(Format.MIND_MAP, SAMPLES[Format.MIND_MAP], DEFAULT, ctx())
    }
    opml = xml.dom.minidom.parseString(files["opml"].content)
    branches = opml.getElementsByTagName("body")[0].firstChild.childNodes
    assert len(branches) == len(SAMPLES[Format.MIND_MAP].branches)
    assert files["mmd"].content.decode().startswith("mindmap")
    assert b"## Diagn" in files["md"].content


def test_concept_map_drawio_has_nodes_and_edges() -> None:
    files = {
        f.file_type: f
        for f in generate_files(Format.CONCEPT_MAP, SAMPLES[Format.CONCEPT_MAP], DEFAULT, ctx())
    }
    doc = xml.dom.minidom.parseString(files["drawio"].content)
    cells = doc.getElementsByTagName("mxCell")
    assert sum(c.getAttribute("vertex") == "1" for c in cells) == 4
    assert sum(c.getAttribute("edge") == "1" for c in cells) == 3
    assert "c_sobregiro" in files["mmd"].content.decode()


def test_commercial_sheet_files() -> None:
    sheet = CommercialSheet.model_validate(
        {
            "product_name": "Adeus Cheque Especial",
            "alternative_names": ["a", "b", "c"],
            "headline": "h",
            "subheadline": "s",
            "currency": "BRL",
            "main_price": {"min": 27, "max": 47},
            "order_bumps": [{"title": "Planilha", "price": 17, "checkout_text": "Adicione"}],
            "bonuses": [{"title": "Mapa", "perceived_value": 47}],
            "short_description": "d",
            "benefit_bullets": ["x"],
            "ad_hooks": ["1", "2", "3"],
        }
    )
    md, docx = commercial_files(sheet, DEFAULT, ctx("pt-BR", slug="ficha-comercial"))
    assert md.filename == "ficha-comercial.md"
    text = md.content.decode()
    assert "R$ 27.00 – R$ 47.00" in text and "Texto do checkout" in text
    assert Document(io.BytesIO(docx.content)).paragraphs


def test_guide_page_count_grows_with_content() -> None:
    base = SAMPLES[Format.PDF_GUIDE]
    bigger = GuideContent.model_validate(
        {**base.model_dump(), "chapters": [c.model_dump() for c in base.chapters] * 3}
    )
    [small] = generate_files(Format.PDF_GUIDE, base, DEFAULT, ctx())
    [large] = generate_files(Format.PDF_GUIDE, bigger, DEFAULT, ctx())
    assert pdf_page_count(large.content) > pdf_page_count(small.content)


def test_slugify() -> None:
    assert slugify("Adiós Sobregiro: ¡60 días!") == "adios-sobregiro-60-dias"
    assert slugify("???") == "producto"
