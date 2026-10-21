"""Fase 7: recorrido completo por la API (como lo hace el frontend), con repos y IA en memoria.

Cubre los criterios de aceptación que se pueden verificar sin servicios externos: el paquete
1 + 3 + 2, formatos distintos, archivos que abren, cambio de plantilla sin IA, regenerar una
pieza, descargas solo para el dueño y los límites de seguridad.
"""

import asyncio
import csv
import io
import xml.etree.ElementTree as ET
import zipfile
from types import SimpleNamespace
from uuid import UUID, uuid4

import openpyxl
import pypdfium2
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import Settings, get_settings
from app.content.engines import ContentService, EngineConfig
from app.content.prompts import PromptLibrary
from app.content.schemas import Format, ProblemProposals
from app.deps import get_services
from app.generators.samples import SAMPLES
from app.main import create_app
from app.worker import run_one
from tests.fakes import FakeLLM, FakeUsage, make_outline, make_problem
from tests.memory import MemoryFiles, MemoryItems, MemoryJobs, MemoryProducts, MemoryStorage
from tests.test_auth import ALLOWED, SECRET, make_token
from tests.test_engines import commercial

SETTINGS = Settings(
    _env_file=None,
    supabase_jwt_secret=SECRET,
    allowed_emails=[ALLOWED],
    rate_limit_per_minute=1000,
)
USER = str(uuid4())


class Products(MemoryProducts):
    """Como ProductRepo real: `/outline` crea el producto en `outline_ready`."""

    async def create_with_outline(self, *, user_id, request, template_id, outline, prompt_version):
        row = self.add(
            user_id=str(user_id),
            title=outline.product_name,
            niche=request.niche,
            problem=request.problem.model_dump(),
            market=request.market,
            main_format=request.main_format.value,
            template_id=template_id,
            outline=outline.model_dump(mode="json"),
            prompt_version=prompt_version,
        )
        return UUID(row["id"])


class App:
    def __init__(self) -> None:
        self.llm = FakeLLM()
        self.products = Products()
        self.items, self.files = MemoryItems(), MemoryFiles()
        self.jobs, self.storage = MemoryJobs(), MemoryStorage()
        prompts = PromptLibrary(SETTINGS.prompts_dir)
        content = ContentService(
            llm=self.llm, prompts=prompts, usage=FakeUsage(), config=EngineConfig("m", "m")
        )
        deps = SimpleNamespace(
            content=content,
            products=self.products,
            items=self.items,
            files=self.files,
            jobs=self.jobs,
            storage=self.storage,
            templates_dir=SETTINGS.templates_dir,
            sync_target=None,
            uploads=None,
        )
        self.services = SimpleNamespace(
            settings=SETTINGS,
            prompts=prompts,
            content=content,
            products=self.products,
            items=self.items,
            files=self.files,
            jobs=self.jobs,
            storage=self.storage,
            pipeline=lambda: deps,
        )
        app = create_app()
        app.dependency_overrides[get_settings] = lambda: SETTINGS
        app.dependency_overrides[get_services] = lambda: self.services
        self.client = TestClient(app)
        self.as_user(USER)

    def as_user(self, user: str) -> None:
        self.client.headers["Authorization"] = f"Bearer {make_token(sub=user)}"

    def work(self) -> None:
        assert asyncio.run(run_one(self.services)) is True

    def files_of(self, product_id: str) -> list[dict]:
        return [f for f in self.files.rows.values() if f["product_id"] == product_id]

    def body(self, file: dict) -> bytes:
        return self.storage.objects[file["storage_path"]]


def outline_samples(outline) -> list:
    return [SAMPLES[Format(i.format)] for i in outline.items()]


@pytest.fixture(scope="module")
def journey():
    """Un único recorrido: problemas → outline → editar → generar. Los tests lo inspeccionan."""
    app = App()
    outline = make_outline()
    app.llm.responses.extend(
        [
            ProblemProposals(problems=[make_problem()] * 5),
            outline,
            *outline_samples(outline),
            commercial("BRL"),
        ]
    )
    c = app.client
    problems = c.post("/problems", json={"niche": "Finanças pessoais", "market": "pt-BR"})
    built = c.post(
        "/outline",
        json={
            "niche": "Finanças pessoais",
            "problem": problems.json()["problems"][0],
            "market": "pt-BR",
            "main_format": "pdf_guide",
            "template_id": "minimal-clean",
        },
    )
    product_id = built.json()["product_id"]
    edited = built.json()["outline"]
    edited["product_name"] = "Adeus Cheque Especial (editado)"
    saved = c.put(f"/products/{product_id}/outline", json=edited)
    generated = c.post(f"/products/{product_id}/generate")
    again = c.post(f"/products/{product_id}/generate")  # mientras el job está en cola
    app.work()
    return SimpleNamespace(
        app=app,
        problems=problems,
        built=built,
        saved=saved,
        generated=generated,
        again=again,
        product_id=product_id,
        calls=len(app.llm.calls),
    )


def test_api_requires_a_valid_session(journey) -> None:
    anonymous = journey.app.client.post(
        "/problems", json={"niche": "x" * 5}, headers={"Authorization": ""}
    )
    assert anonymous.status_code == 401
    stranger = make_token(email="otra@example.com", sub=str(uuid4()))
    forbidden = journey.app.client.get("/me", headers={"Authorization": f"Bearer {stranger}"})
    assert forbidden.status_code == 403


def test_wizard_steps_return_what_the_ui_needs(journey) -> None:
    assert journey.problems.status_code == 200 and len(journey.problems.json()["problems"]) == 5
    assert journey.built.status_code == 200
    assert journey.saved.status_code == 200
    assert journey.generated.status_code == 202
    product = journey.app.products.rows[journey.product_id]
    assert product["title"] == "Adeus Cheque Especial (editado)"
    assert product["status"] == "ready"


def test_a_second_generation_is_refused_while_one_is_running(journey) -> None:
    assert journey.again.status_code == 409


def test_package_is_one_main_three_bumps_two_bonuses_with_three_formats(journey) -> None:
    items = list(journey.app.items.rows.values())
    roles = sorted(i["role"] for i in items)
    assert roles == ["bonus"] * 2 + ["main"] + ["order_bump"] * 3
    assert all(i["status"] == "ready" for i in items)
    assert len({i["format"] for i in items}) >= 3


def test_every_generated_file_opens(journey) -> None:
    app = journey.app
    files = [f for f in app.files_of(journey.product_id) if f["file_type"] != "zip"]
    seen = set()
    for f in files:
        data, kind = app.body(f), f["file_type"]
        seen.add(kind)
        assert data, f"{kind} vacío"
        if kind == "pdf":
            assert len(pypdfium2.PdfDocument(data)) >= 1
        elif kind == "docx":
            assert Document(io.BytesIO(data)).paragraphs
        elif kind == "xlsx":
            assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames
        elif kind == "csv":
            assert len(list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))) >= 2
        elif kind in ("svg", "drawio", "opml", "mm"):
            ET.fromstring(data)
        elif kind == "png":
            assert data.startswith(b"\x89PNG")
        elif kind == "xmind":
            assert zipfile.ZipFile(io.BytesIO(data)).namelist()
        elif kind in ("md", "mmd"):
            assert data.decode("utf-8").strip()
    assert {"pdf", "xlsx", "docx", "csv", "svg", "xmind", "md"} <= seen


def test_spreadsheets_keep_working_formulas(journey) -> None:
    app = journey.app
    sheet = next(f for f in app.files_of(journey.product_id) if f["file_type"] == "xlsx")
    book = openpyxl.load_workbook(io.BytesIO(app.body(sheet)))  # sin data_only: ve las fórmulas
    formulas = [
        c.value
        for ws in book
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert formulas, "el XLSX no tiene fórmulas"


def test_zip_contains_every_file_and_the_commercial_sheet(journey) -> None:
    app = journey.app
    files = app.files_of(journey.product_id)
    zip_row = next(f for f in files if f["file_type"] == "zip")
    names = zipfile.ZipFile(io.BytesIO(app.body(zip_row))).namelist()
    assert "ficha-comercial.md" in names
    assert {n.split("/")[0] for n in names if "/" in n} == {"principal", "order-bumps", "bonus"}
    assert len(names) >= len(files) - 1  # todos menos el propio ZIP
    assert app.products.rows[journey.product_id]["commercial"]["currency"] == "BRL"


def test_downloads_are_signed_and_only_for_the_owner(journey) -> None:
    app = journey.app
    file = app.files_of(journey.product_id)[0]
    ok = app.client.get(f"/files/{file['id']}/download")
    assert ok.status_code == 200 and ok.json()["url"]
    app.as_user(str(uuid4()))
    try:
        assert app.client.get(f"/files/{file['id']}/download").status_code == 404
        assert app.client.post(f"/products/{journey.product_id}/generate").status_code == 404
    finally:
        app.as_user(USER)


def test_template_change_creates_a_linked_version_without_calling_the_ai(journey) -> None:
    app = journey.app
    before = len(app.llm.calls)
    response = app.client.post(
        f"/products/{journey.product_id}/remodel", json={"template_id": "dark-elegante"}
    )
    assert response.status_code == 202
    assert response.json()["plan"] == {"generate": 0, "rerender": 6}
    app.work()
    assert len(app.llm.calls) == before
    child = app.products.rows[response.json()["product_id"]]
    assert child["status"] == "ready" and child["parent_product_id"] == journey.product_id
    assert {f["file_type"] for f in app.files_of(child["id"])} >= {"pdf", "xlsx", "zip"}


def test_regenerating_one_piece_calls_the_ai_once(journey) -> None:
    app = journey.app
    bonus = next(i for i in app.items.rows.values() if i["role"] == "bonus")
    app.llm.responses.append(SAMPLES[Format(bonus["format"])])
    before = len(app.llm.calls)
    response = app.client.post(f"/products/{journey.product_id}/items/{bonus['id']}/regenerate")
    assert response.status_code == 202
    app.work()
    assert len(app.llm.calls) == before + 1
    assert app.products.rows[journey.product_id]["status"] == "ready"


def test_security_headers_and_rate_limit_on_the_real_app(journey, monkeypatch) -> None:
    response = journey.app.client.get("/health")
    assert response.headers["x-content-type-options"] == "nosniff"
    strict = Settings(
        _env_file=None,
        supabase_jwt_secret=SECRET,
        allowed_emails=[ALLOWED],
        rate_limit_per_minute=2,
    )
    journey.app.client.app.dependency_overrides[get_settings] = lambda: strict
    try:
        url = f"/products/{journey.product_id}/generate"
        statuses = [journey.app.client.post(url).status_code for _ in range(4)]
        assert statuses[-1] == 429
    finally:
        journey.app.client.app.dependency_overrides[get_settings] = lambda: SETTINGS
